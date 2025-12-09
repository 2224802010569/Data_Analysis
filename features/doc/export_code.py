import os
from docx import Document

VERSION = "1.3.0"
SERVICE_NAME = "Back_end"
OUTPUT_NAME = f"{VERSION}_{SERVICE_NAME}.docx"

def export_code():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    service_path = os.path.abspath(os.path.join(current_dir, ".."))
    print("[DEBUG] Đường dẫn đến service:", service_path)

    if not os.path.exists(service_path):
        print("[❌] Không tìm thấy thư mục:", service_path)
        return

    doc = Document()
    doc.add_heading(f"📦 Mã nguồn: {SERVICE_NAME}", level=1)
    file_count = 0
    for root, dirs, files in os.walk(service_path):
        print("[DEBUG] Đang đọc thư mục:", root)
        for file in files:
            if not file.endswith((".py", ".html", ".css", ".js")):
                continue
            file_path = os.path.join(root, file)
            rel_path = os.path.relpath(file_path, service_path)
            print(f"📄 Đọc file: {rel_path}")
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                doc.add_heading(rel_path, level=2)
                doc.add_paragraph(content, style='Code')
                file_count += 1
            except Exception as e:
                print(f"Không đọc được file: {file_path} → {type(e).__name__}: {str(e)}")
    output_path = os.path.abspath(os.path.join(current_dir, "../doc", OUTPUT_NAME))
    doc.save(output_path)
    print(f"✅ Hoàn tất. Đã ghi {file_count} file vào: {output_path}")

if __name__ == '__main__':
    export_code()
